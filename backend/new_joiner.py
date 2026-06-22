"""
New-joiner intake API.

Receives a multi-step form submission from /newjoinee, organises files into
a fixed folder structure on disk, generates a Word doc of the text fields,
and notifies HR via Resend.

Storage root is configurable via NEW_JOINER_STORAGE_ROOT env var (default
/app/data/new-joiners — fine for the K8s preview; Hetzner mounts a persistent
path here in production).
"""

from __future__ import annotations

import asyncio
import json
import logging
import os
import re
import unicodedata
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

import resend
from docx import Document
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, Form, HTTPException, Request, UploadFile
from starlette.datastructures import UploadFile as StarletteUploadFile

logger = logging.getLogger("origami.new_joiner")

router = APIRouter()

STORAGE_ROOT = Path(os.environ.get("NEW_JOINER_STORAGE_ROOT", "/app/data/new-joiners"))
HR_EMAIL = os.environ.get("HR_EMAIL", "hr@origami.dev")
SENDER_EMAIL = os.environ.get("SENDER_EMAIL", "onboarding@resend.dev")

# 10 MB per file, 250 MB per submission (defensive ceiling).
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TOTAL_BYTES = 250 * 1024 * 1024
ALLOWED_EXT = {".pdf", ".jpg", ".jpeg", ".png"}

# Map form-field name → (subfolder, output filename stem). Multi-file fields
# end with "[]" and write filename_1, filename_2, ...
FILE_PLAN: dict[str, tuple[str, str]] = {
    "marks_10th": ("Educational Certificates", "10th_marks_card"),
    "marks_12th": ("Educational Certificates", "12th_marks_card"),
    "degree_marks": ("Educational Certificates", "degree_marks_card"),
    "pan_card": ("Identity Documents", "pan_card"),
    "aadhaar_card": ("Identity Documents", "aadhaar_card"),
    "passport_copy": ("Identity Documents", "passport_copy"),
    "experience_1": ("Employment History", "experience_letter_employer1"),
    "relieving_1": ("Employment History", "relieving_letter_employer1"),
    "experience_2": ("Employment History", "experience_letter_employer2"),
    "relieving_2": ("Employment History", "relieving_letter_employer2"),
    "experience_3": ("Employment History", "experience_letter_employer3"),
    "relieving_3": ("Employment History", "relieving_letter_employer3"),
    "payslip_1": ("Payslips", "payslip_month1"),
    "payslip_2": ("Payslips", "payslip_month2"),
    "payslip_3": ("Payslips", "payslip_month3"),
    "address_proof": ("Address Proof", "address_proof"),
    "photos[]": ("Photos", "passport_photo"),
}

REQUIRED_FILES = {
    "marks_10th",
    "marks_12th",
    "degree_marks",
    "pan_card",
    "aadhaar_card",
    "experience_1",
    "relieving_1",
    "payslip_1",
    "payslip_2",
    "payslip_3",
    "address_proof",
    "photos[]",
}

REQUIRED_TEXT_FIELDS = (
    "full_name email date_of_joining department mobile_number "
    "bank_name account_holder account_number ifsc_code uan_number "
    "current_address permanent_address "
    "emergency_name emergency_phone emergency_relationship "
    "pro_ref1_name pro_ref1_company pro_ref1_designation pro_ref1_phone pro_ref1_email "
    "pro_ref2_name pro_ref2_company pro_ref2_designation pro_ref2_phone pro_ref2_email "
    "per_ref1_name per_ref1_relationship per_ref1_phone per_ref1_email "
    "per_ref2_name per_ref2_relationship per_ref2_phone per_ref2_email"
).split()


def _slugify(s: str) -> str:
    """Filesystem-safe slug — keeps unicode names readable but strips path chars."""
    s = unicodedata.normalize("NFKC", s).strip()
    s = re.sub(r"[\\/\x00-\x1f<>:\"|?*]+", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s or "unnamed"


def _safe_ext(filename: str | None) -> str:
    if not filename:
        return ".bin"
    ext = Path(filename).suffix.lower()
    return ext if ext in ALLOWED_EXT else ".bin"


def _build_docx(data: dict[str, Any]) -> bytes:
    doc = Document()
    title = doc.add_heading(f"{data['full_name']} — New Joiner Details", level=0)
    for r in title.runs:
        r.font.color.rgb = RGBColor(0xE8, 0x38, 0x2C)

    doc.add_paragraph(
        f"Submitted {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    ).italic = True

    def section(title_: str, rows: list[tuple[str, str]]) -> None:
        doc.add_heading(title_, level=1)
        t = doc.add_table(rows=len(rows), cols=2)
        t.style = "Light Grid Accent 1"
        for i, (k, v) in enumerate(rows):
            cells = t.rows[i].cells
            cells[0].text = k
            cells[1].text = str(v or "")
            for run in cells[0].paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(10)
        doc.add_paragraph()

    section("Personal", [
        ("Full name", data.get("full_name")),
        ("Personal email", data.get("email")),
        ("Date of joining", data.get("date_of_joining")),
        ("Department", data.get("department")),
        ("Mobile number", data.get("mobile_number")),
    ])

    section("Bank details", [
        ("Bank name", data.get("bank_name")),
        ("Account holder name", data.get("account_holder")),
        ("Account number", data.get("account_number")),
        ("IFSC code", data.get("ifsc_code")),
    ])

    section("Provident fund", [
        ("UAN number", data.get("uan_number")),
    ])

    section("Addresses", [
        ("Current address", data.get("current_address")),
        ("Permanent address", data.get("permanent_address")),
        ("Same as current", "Yes" if data.get("permanent_same_as_current") in (True, "true", "on", "1") else "No"),
    ])

    section("Emergency contact", [
        ("Name", data.get("emergency_name")),
        ("Phone", data.get("emergency_phone")),
        ("Relationship", data.get("emergency_relationship")),
    ])

    section("Professional reference 1", [
        ("Name", data.get("pro_ref1_name")),
        ("Company", data.get("pro_ref1_company")),
        ("Designation", data.get("pro_ref1_designation")),
        ("Phone", data.get("pro_ref1_phone")),
        ("Email", data.get("pro_ref1_email")),
    ])

    section("Professional reference 2", [
        ("Name", data.get("pro_ref2_name")),
        ("Company", data.get("pro_ref2_company")),
        ("Designation", data.get("pro_ref2_designation")),
        ("Phone", data.get("pro_ref2_phone")),
        ("Email", data.get("pro_ref2_email")),
    ])

    section("Personal reference 1", [
        ("Name", data.get("per_ref1_name")),
        ("Relationship", data.get("per_ref1_relationship")),
        ("Phone", data.get("per_ref1_phone")),
        ("Email", data.get("per_ref1_email")),
    ])

    section("Personal reference 2", [
        ("Name", data.get("per_ref2_name")),
        ("Relationship", data.get("per_ref2_relationship")),
        ("Phone", data.get("per_ref2_phone")),
        ("Email", data.get("per_ref2_email")),
    ])

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_email_html(data: dict[str, Any], folder_rel: str, file_count: int) -> str:
    rows = "".join(
        f'<tr><td style="padding:6px 12px;color:#6B6B6B;font-size:13px;text-transform:uppercase;letter-spacing:0.08em;width:160px;">{k}</td>'
        f'<td style="padding:6px 12px;color:#0D0D0D;font-size:14px;">{v or "—"}</td></tr>'
        for k, v in [
            ("Name", data.get("full_name")),
            ("Email", data.get("email")),
            ("Mobile", data.get("mobile_number")),
            ("Department", data.get("department")),
            ("Joining date", data.get("date_of_joining")),
        ]
    )
    return f"""
<!DOCTYPE html><html><body style="margin:0;padding:0;background:#FAFAF7;font-family:Inter,Arial,sans-serif;">
  <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#FAFAF7;padding:32px 0;">
    <tr><td align="center">
      <table role="presentation" width="640" cellpadding="0" cellspacing="0" style="background:#FFFFFF;border:1px solid rgba(13,13,13,0.08);">
        <tr><td style="padding:28px 32px 12px;border-bottom:1px solid rgba(13,13,13,0.08);">
          <div style="font-family:'Courier New',monospace;font-size:11px;letter-spacing:0.25em;color:#E8382C;text-transform:uppercase;">/// New joiner submission</div>
          <h1 style="margin:8px 0 0;font-size:26px;font-weight:900;color:#0D0D0D;letter-spacing:-0.02em;">{data.get('full_name')}</h1>
        </td></tr>
        <tr><td style="padding:18px 20px;">
          <table role="presentation" width="100%" cellpadding="0" cellspacing="0">{rows}</table>
        </td></tr>
        <tr><td style="padding:18px 32px;border-top:1px solid rgba(13,13,13,0.08);background:#FAFAF7;">
          <div style="color:#6B6B6B;font-size:13px;line-height:1.7;">
            Files stored on Hetzner at:<br>
            <code style="font-family:'Courier New',monospace;color:#0D0D0D;font-size:13px;">{folder_rel}</code><br>
            {file_count} file(s) uploaded · 1 details.docx generated
          </div>
        </td></tr>
      </table>
    </td></tr></table></body></html>""".strip()


async def _send_hr_email(data: dict[str, Any], folder_rel: str, file_count: int) -> str | None:
    if not resend.api_key:
        logger.error("RESEND_API_KEY missing — HR email not sent")
        return None
    params = {
        "from": SENDER_EMAIL,
        "to": [HR_EMAIL],
        "reply_to": data.get("email") or HR_EMAIL,
        "subject": f"New joiner submission — {data['full_name']}",
        "html": _build_email_html(data, folder_rel, file_count),
    }
    try:
        email = await asyncio.to_thread(resend.Emails.send, params)
        return email.get("id") if isinstance(email, dict) else None
    except Exception as exc:  # noqa: BLE001
        logger.exception("Failed to send HR notification: %s", exc)
        return None


@router.post("/api/newjoinee")
async def submit_new_joiner(request: Request):
    """
    Multipart submission. Field names listed in REQUIRED_TEXT_FIELDS +
    FILE_PLAN. Multi-file photos field is `photos[]`.
    """
    form = await request.form()

    # Validate text fields
    data: dict[str, Any] = {}
    for f in REQUIRED_TEXT_FIELDS:
        v = form.get(f)
        if v is None or (isinstance(v, str) and v.strip() == ""):
            raise HTTPException(status_code=422, detail=f"Missing required field: {f}")
        data[f] = v.strip() if isinstance(v, str) else v

    # Optional text fields
    data["permanent_same_as_current"] = form.get("permanent_same_as_current")

    # Build the per-joiner folder
    full_name = _slugify(data["full_name"])
    folder_name = f"{full_name} - {datetime.now(timezone.utc).strftime('%Y-%m-%d')}"
    root_rel = Path("HR/New Joiners") / folder_name
    root_abs = STORAGE_ROOT / root_rel
    if root_abs.exists():
        # Avoid clobbering an earlier submission with the same name today
        suffix = datetime.now(timezone.utc).strftime('%H%M%S')
        root_abs = root_abs.with_name(f"{folder_name} ({suffix})")
        root_rel = root_rel.with_name(root_abs.name)
    root_abs.mkdir(parents=True, exist_ok=False)

    # Pre-create subfolders
    for subdir in {v[0] for v in FILE_PLAN.values()}:
        (root_abs / subdir).mkdir(exist_ok=True)

    saved: list[str] = []
    total_bytes = 0
    received_fields: set[str] = set()

    # Iterate multi-part entries (form.multi_items() preserves duplicates for [])
    for key, value in form.multi_items():
        if not isinstance(value, (UploadFile, StarletteUploadFile)):
            continue
        plan = FILE_PLAN.get(key)
        if plan is None:
            logger.warning("Ignoring unknown file field: %s", key)
            continue
        subdir, stem = plan
        ext = _safe_ext(value.filename)
        if ext == ".bin":
            raise HTTPException(status_code=415, detail=f"{value.filename}: only PDF/JPG/PNG accepted.")

        await value.seek(0)
        contents = await value.read()
        size = len(contents)
        if size == 0:
            logger.warning("Empty upload for field=%s filename=%s", key, value.filename)
            continue
        if size > MAX_FILE_BYTES:
            raise HTTPException(status_code=413, detail=f"{value.filename} exceeds 10 MB limit.")
        total_bytes += size
        if total_bytes > MAX_TOTAL_BYTES:
            raise HTTPException(status_code=413, detail="Total upload size exceeds 250 MB.")

        # Compose target filename (with index for multi-file photos)
        if key.endswith("[]"):
            existing = len([p for p in (root_abs / subdir).iterdir() if p.stem.startswith(stem)])
            target = root_abs / subdir / f"{stem}_{existing + 1}{ext}"
        else:
            target = root_abs / subdir / f"{stem}{ext}"

        target.write_bytes(contents)
        saved.append(str(target.relative_to(STORAGE_ROOT)))
        received_fields.add(key)

    # Verify required files are present
    missing = REQUIRED_FILES - received_fields
    if missing:
        # Cleanup the partial folder so we don't pollute storage
        try:
            for p in root_abs.rglob("*"):
                if p.is_file():
                    p.unlink()
            for p in sorted(root_abs.rglob("*"), reverse=True):
                if p.is_dir():
                    p.rmdir()
            root_abs.rmdir()
        except Exception:  # noqa: BLE001
            logger.exception("Cleanup after validation failure failed")
        raise HTTPException(
            status_code=422,
            detail=f"Missing required files: {', '.join(sorted(missing))}",
        )

    # Generate Details.docx
    docx_bytes = _build_docx(data)
    docx_path = root_abs / f"{full_name} - Details.docx"
    docx_path.write_bytes(docx_bytes)

    # Persist a small JSON manifest for HR ops / debugging
    manifest = {
        "submitted_at": datetime.now(timezone.utc).isoformat(),
        "data": data,
        "files": saved,
        "docx": str(docx_path.relative_to(STORAGE_ROOT)),
    }
    (root_abs / "_manifest.json").write_text(json.dumps(manifest, indent=2, default=str))

    # Notify HR
    email_id = await _send_hr_email(data, str(root_rel), len(saved))

    logger.info(
        "new_joiner saved name=%s folder=%s files=%d email_id=%s",
        data["full_name"], root_rel, len(saved), email_id,
    )

    return {
        "status": "success",
        "folder": str(root_rel),
        "files_saved": len(saved),
        "email_id": email_id,
    }
